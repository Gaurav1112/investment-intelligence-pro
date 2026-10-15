"""Holdings CSV/XLSX import hardening (threat-model C-07, SEC-30..41).

Pipeline here: limits → type check → decode → parse → validate → **preview**. Applying a preview (one transaction
plus the import ledger) belongs to the store layer; ``confirm_preview`` only checks that the user confirmed the
exact file that was previewed.

Holdings template (no column spec exists in the product docs — see the report's spec issues):
``instrument_name, isin, amfi_code, units, cost_basis_inr, purchase_date``; ``instrument_name`` and ``units`` are
required, unknown columns are ignored with a warning.
"""

from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook

from iip.data.export import TRIGGERS
from iip.data.models import DataError, RowRejection, sha256_hex
from iip.data.validators import validate_isin

MiB = 1024 * 1024
MAX_IMPORT_BYTES = 5 * MiB
MAX_ROWS = 20_000
MAX_COLUMNS = 64
MAX_CELL_CHARS = 512
MAX_LINE_BYTES = 1 * MiB
MAX_LABEL_CHARS = 120
XLSX_MAX_MEMBERS = 100
XLSX_MAX_UNCOMPRESSED = 50 * MiB
XLSX_MAX_RATIO = 100

COLUMNS = ("instrument_name", "isin", "amfi_code", "units", "cost_basis_inr", "purchase_date")
REQUIRED_COLUMNS = frozenset({"instrument_name", "units"})
_ENCODINGS = {"utf-8": "utf-8-sig", "cp1252": "cp1252", "utf-16": "utf-16"}
_REJECTED_EXT = frozenset({".xlsm", ".xls", ".xlsb", ".ods", ".xltm", ".xlam"})
_ZIP_MAGIC = b"PK\x03\x04"

_PLAIN = r"[0-9]+(?:\.[0-9]+)?"
_WESTERN = r"[0-9]{1,3}(?:,[0-9]{3})+(?:\.[0-9]+)?"
_INDIAN = r"[0-9]{1,2}(?:,[0-9]{2})*,[0-9]{3}(?:\.[0-9]+)?"
_AMOUNT_RE = re.compile(rf"-?(?:{_PLAIN}|{_WESTERN}|{_INDIAN})")
_CODE_RE = re.compile(r"[0-9]{1,9}")
_ISO_DATE_RE = re.compile(r"([0-9]{4})-([0-9]{2})-([0-9]{2})")
_DMY_RE = re.compile(r"([0-9]{2})-([0-9]{2})-([0-9]{4})")
_DMONY_RE = re.compile(r"([0-9]{2})-([A-Z][a-z]{2})-([0-9]{4})")
_MONTHS = {
    m: i
    for i, m in enumerate(("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"), start=1)
}
_BIDI = frozenset(
    [chr(c) for c in range(0x202A, 0x202F)] + [chr(c) for c in range(0x2066, 0x206A)] + ["\u200e", "\u200f", "\u061c"]
)


class ImportRejected(DataError):
    """The whole file is refused. ``status`` is the HTTP status the API should return (413/415/422/409)."""

    def __init__(self, code: str, message: str, status: int = 422) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status = status


@dataclass(frozen=True, slots=True)
class HoldingRow:
    instrument_name: str
    isin: str | None
    amfi_code: int | None
    units: Decimal
    cost_basis_inr: Decimal | None
    purchase_date: date | None
    flags: tuple[str, ...] = ()

    def as_export_dict(self) -> dict[str, object]:
        return {
            "instrument_name": self.instrument_name,
            "isin": self.isin,
            "amfi_code": self.amfi_code,
            "units": self.units,
            "cost_basis_inr": self.cost_basis_inr,
            "purchase_date": self.purchase_date,
        }


@dataclass(frozen=True, slots=True)
class ImportPreview:
    kind: str  # "csv" | "xlsx"
    label: str
    file_sha256: str
    header: tuple[str, ...]
    rows: tuple[HoldingRow, ...]
    rejected: tuple[RowRejection, ...]
    warnings: tuple[str, ...]


# --- text and number helpers ----------------------------------------------------------------------------------------


def normalise_text(text: str) -> str:
    """NFC; strip control characters (tab kept inside the value) and bidi overrides; trim."""
    text = unicodedata.normalize("NFC", text)
    text = "".join(ch for ch in text if ch not in _BIDI and (ch == "\t" or unicodedata.category(ch) != "Cc"))
    return text.strip()


def unescape_export_prefix(text: str) -> str:
    """Reverse our own export escaping: ``'`` followed by a trigger character."""
    if len(text) > 1 and text[0] == "'" and text[1] in TRIGGERS:
        return text[1:]
    return text


def sanitise_label(filename: str) -> str:
    """Display label for an uploaded file: basename only, cleaned, ≤ 120 characters."""
    cleaned = normalise_text(filename.replace("\\", "/"))
    name = PurePosixPath(cleaned).name if cleaned else ""
    name = name.lstrip(".") or "upload"
    if len(name) > MAX_LABEL_CHARS:
        suffix = PurePosixPath(name).suffix[:10]
        name = name[: MAX_LABEL_CHARS - len(suffix)] + suffix
    return name


def parse_amount(text: str) -> Decimal:
    """Strict amount parser: plain, Western (``100,000.50``) or Indian (``1,00,000.50``) grouping; optional
    leading ``-``. Rejects ``1.000,50``, exponents, currency symbols, spaces and non-ASCII digits."""
    s = text.strip()
    if not _AMOUNT_RE.fullmatch(s):
        raise ValueError(f"not an accepted number format: {text!r}")
    return Decimal(s.replace(",", ""))


def parse_import_date(text: str) -> date:
    """ISO ``YYYY-MM-DD``, ``DD-MM-YYYY`` or ``DD-Mon-YYYY``. Slash formats are ambiguous and rejected."""
    s = text.strip()
    if m := _ISO_DATE_RE.fullmatch(s):
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    if m := _DMY_RE.fullmatch(s):
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    if (m := _DMONY_RE.fullmatch(s)) and m.group(2) in _MONTHS:
        return date(int(m.group(3)), _MONTHS[m.group(2)], int(m.group(1)))
    raise ValueError(f"unsupported date format: {text!r}")


# --- entry points ---------------------------------------------------------------------------------------------------


def preview_import(data: bytes, filename: str, *, encoding: str | None = None) -> ImportPreview:
    """Validate an uploaded holdings file and return a preview. Raises ``ImportRejected`` for file-level problems;
    row-level problems are listed in ``ImportPreview.rejected``."""
    if len(data) > MAX_IMPORT_BYTES:
        raise ImportRejected("TOO_LARGE", f"file exceeds {MAX_IMPORT_BYTES} bytes", status=413)
    label = sanitise_label(filename)
    ext = PurePosixPath(label).suffix.lower()
    if ext in _REJECTED_EXT:
        raise ImportRejected("UNSUPPORTED_TYPE", f"{ext} files are not accepted (no macros, no legacy formats)", 415)
    is_zip = data[:4] == _ZIP_MAGIC
    warnings: list[str] = []
    if ext == ".csv":
        if is_zip:
            raise ImportRejected("TYPE_MISMATCH", "file named .csv is a ZIP container", 415)
        text = _decode(data, encoding)
        table = _parse_csv_text(text)
        kind = "csv"
    elif ext == ".xlsx":
        if not is_zip:
            raise ImportRejected("TYPE_MISMATCH", "file named .xlsx is not a ZIP container", 415)
        _inspect_xlsx(data)
        table = _read_xlsx(data, warnings)
        kind = "xlsx"
    else:
        raise ImportRejected("UNSUPPORTED_TYPE", "only .csv and .xlsx files are accepted", 415)
    header, rows, rejected = _build_rows(table, warnings)
    return ImportPreview(kind, label, sha256_hex(data), header, rows, rejected, tuple(warnings))


def confirm_preview(preview: ImportPreview, confirmed_sha256: str) -> ImportPreview:
    """The user must confirm the preview of exactly this file before it is applied."""
    if confirmed_sha256 != preview.file_sha256:
        raise ImportRejected("PREVIEW_MISMATCH", "confirmation does not match the previewed file", 409)
    return preview


# --- CSV --------------------------------------------------------------------------------------------------------------

Table = list[tuple[int, list[Any]]]


def _decode(data: bytes, encoding: str | None) -> str:
    if encoding is None:
        try:
            if b"\x00" in data:
                raise UnicodeDecodeError("utf-8", data, 0, 1, "NUL byte")
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ImportRejected(
                "ENCODING_NOT_UTF8",
                "The file is not valid UTF-8. Choose its encoding explicitly (cp1252 or utf-16) and preview again.",
            ) from None
    codec = _ENCODINGS.get(encoding.lower())
    if codec is None:
        raise ImportRejected("ENCODING_UNSUPPORTED", f"encoding must be one of {sorted(_ENCODINGS)}")
    try:
        text = data.decode(codec)
    except UnicodeDecodeError:
        raise ImportRejected("ENCODING_ERROR", f"the file does not decode as {encoding}") from None
    if "\x00" in text:
        raise ImportRejected("NUL_BYTE", "the file contains NUL characters")
    return text.removeprefix("\ufeff")


def _parse_csv_text(text: str) -> Table:
    for i, line in enumerate(text.splitlines(), start=1):
        if len(line.encode("utf-8")) > MAX_LINE_BYTES:
            raise ImportRejected("LINE_TOO_LONG", f"row {i} exceeds {MAX_LINE_BYTES} bytes")
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    table: Table = []
    try:
        for cells in reader:
            line_no = reader.line_num
            if not cells or all(c.strip() == "" for c in cells):
                continue
            _check_shape(line_no, cells, len(table))
            table.append((line_no, list(cells)))
    except csv.Error as exc:
        raise ImportRejected("CSV_MALFORMED", f"row {reader.line_num}: {exc}") from None
    return table


def _check_shape(line_no: int, cells: list[Any], rows_so_far: int) -> None:
    if rows_so_far > MAX_ROWS:  # rows_so_far includes the header row
        raise ImportRejected("TOO_MANY_ROWS", f"row {line_no}: more than {MAX_ROWS} data rows")
    if len(cells) > MAX_COLUMNS:
        raise ImportRejected("TOO_MANY_COLUMNS", f"row {line_no}: more than {MAX_COLUMNS} columns")
    for c in cells:
        if isinstance(c, str) and len(c) > MAX_CELL_CHARS:
            raise ImportRejected("CELL_TOO_LONG", f"row {line_no}: a cell exceeds {MAX_CELL_CHARS} characters")


# --- XLSX ------------------------------------------------------------------------------------------------------------


def _inspect_xlsx(data: bytes) -> None:
    """Container checks before openpyxl sees the bytes (zip bomb, macros, links, OLE, DTDs, paths)."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ImportRejected("NOT_XLSX", "not a valid XLSX (ZIP) file", 415) from None
    with zf:
        infos = zf.infolist()
        if len(infos) > XLSX_MAX_MEMBERS:
            raise ImportRejected("XLSX_TOO_MANY_MEMBERS", f"more than {XLSX_MAX_MEMBERS} archive members")
        names = {i.filename for i in infos}
        if "[Content_Types].xml" not in names:
            raise ImportRejected("NOT_XLSX", "missing [Content_Types].xml", 415)
        total = 0
        for info in infos:
            name = info.filename
            low = name.lower()
            parts = name.split("/")
            if name.startswith("/") or "\\" in name or ".." in parts or ":" in name:
                raise ImportRejected("XLSX_BAD_PATH", "archive member with an absolute or parent path")
            if "vbaproject" in low:
                raise ImportRejected("XLSX_MACRO", "workbook contains a VBA project (macros)")
            if low.startswith("xl/externallinks/"):
                raise ImportRejected("XLSX_EXTERNAL_LINK", "workbook contains external links")
            if "embeddings/" in low or "oleobject" in low or "activex" in low:
                raise ImportRejected("XLSX_OLE", "workbook contains embedded OLE/ActiveX objects")
            if info.file_size > 0 and info.file_size > XLSX_MAX_RATIO * max(info.compress_size, 1):
                raise ImportRejected("XLSX_RATIO", f"member compression ratio exceeds {XLSX_MAX_RATIO}:1")
            total += info.file_size
            if total > XLSX_MAX_UNCOMPRESSED:
                raise ImportRejected("XLSX_TOO_LARGE", "uncompressed size exceeds 50 MiB")
        content_types = zf.read("[Content_Types].xml")
        if b"macroenabled" in content_types.lower():
            raise ImportRejected("XLSX_MACRO", "content types declare a macro-enabled workbook")
        for info in infos:
            if info.filename.lower().endswith((".xml", ".rels", ".vml")):
                head = zf.read(info).lower()
                if b"<!doctype" in head or b"<!entity" in head:
                    raise ImportRejected("XLSX_DTD", "XML with a DOCTYPE/ENTITY declaration is refused")


def _load_workbook(data: bytes, *, data_only: bool) -> Any:
    return load_workbook(io.BytesIO(data), read_only=True, data_only=data_only, keep_links=False, keep_vba=False)


def _read_xlsx(data: bytes, warnings: list[str]) -> Table:
    """Read the first sheet. Formulas are never evaluated: one pass sees formula text, the other the cached values.
    A formula with no cached value becomes empty with a warning."""
    try:
        wb_f = _load_workbook(data, data_only=False)
        wb_v = _load_workbook(data, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise ImportRejected("XLSX_UNREADABLE", f"workbook could not be read ({exc.__class__.__name__})") from None
    try:
        ws_f, ws_v = wb_f.worksheets[0], wb_v.worksheets[0]
        bounds = {"max_row": MAX_ROWS + 2, "max_col": MAX_COLUMNS + 1, "values_only": True}
        table: Table = []
        for row_no, (rf, rv) in enumerate(zip(ws_f.iter_rows(**bounds), ws_v.iter_rows(**bounds), strict=False), 1):
            cells: list[Any] = []
            for col, (f, v) in enumerate(zip(rf, rv, strict=False), start=1):
                if isinstance(f, str) and f.startswith("=") and v != f:
                    if v is None:
                        warnings.append(f"row {row_no} col {col}: FORMULA_WITHOUT_CACHED_VALUE treated as empty")
                    else:
                        warnings.append(f"row {row_no} col {col}: FORMULA_CACHED_VALUE_USED")
                cells.append(_xlsx_cell(v))
            while cells and cells[-1] is None:
                cells.pop()
            if not cells or all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
                continue
            _check_shape(row_no, cells, len(table))
            table.append((row_no, cells))
        return table
    finally:
        wb_f.close()
        wb_v.close()


def _xlsx_cell(v: object) -> object:
    if v is None or (isinstance(v, (str, int, Decimal)) and not isinstance(v, bool)):
        return v
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, float):
        if not math.isfinite(v):
            return "NaN"
        return Decimal(repr(v))  # shortest round-trip repr, never Decimal(float)
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return str(v)


# --- rows ------------------------------------------------------------------------------------------------------------


def _build_rows(
    table: Table, warnings: list[str]
) -> tuple[tuple[str, ...], tuple[HoldingRow, ...], tuple[RowRejection, ...]]:
    if not table:
        raise ImportRejected("EMPTY", "the file has no header row")
    _, raw_header = table[0]
    header = tuple(normalise_text(str(h or "")).lower() for h in raw_header)
    if len(set(header)) != len(header):
        raise ImportRejected("DUPLICATE_COLUMN", "the header repeats a column name")
    missing = REQUIRED_COLUMNS - set(header)
    if missing:
        raise ImportRejected("MISSING_COLUMN", f"missing required column(s): {', '.join(sorted(missing))}")
    for h in header:
        if h not in COLUMNS:
            warnings.append(f"UNKNOWN_COLUMN ignored: {h[:40]}")
    index = {h: i for i, h in enumerate(header) if h in COLUMNS}

    rows: list[HoldingRow] = []
    rejected: list[RowRejection] = []
    for line_no, cells in table[1:]:
        values = {name: (cells[i] if i < len(cells) else None) for name, i in index.items()}
        result = _build_row(values)
        if isinstance(result, str):
            digest = sha256_hex(repr(cells).encode("utf-8"))
            rejected.append(RowRejection(line_no, result, f"row {line_no}", digest))
        else:
            rows.append(result)
    return header, tuple(rows), tuple(rejected)


def _text(v: object) -> str | None:
    if v is None:
        return None
    if isinstance(v, str):
        t = normalise_text(unescape_export_prefix(v))
        return t or None
    return None


def _decimal(v: object) -> Decimal | None:
    """Raises ValueError for anything that is not a strictly formatted number."""
    if v is None:
        return None
    if isinstance(v, Decimal):
        if not v.is_finite():
            raise ValueError("non-finite")
        return v
    if isinstance(v, int) and not isinstance(v, bool):
        return Decimal(v)
    if isinstance(v, str):
        t = normalise_text(v)
        return parse_amount(t) if t else None
    raise ValueError("not a number")


def _build_row(values: dict[str, object]) -> HoldingRow | str:
    flags: list[str] = []
    raw_name = values.get("instrument_name")
    name = _text(raw_name) if not isinstance(raw_name, (int, Decimal)) else str(raw_name)
    if not name:
        return "MISSING_NAME"
    if name[0] in TRIGGERS:
        flags.append("FORMULA_LIKE_TEXT")

    isin = _text(values.get("isin"))
    if isin is not None and (reason := validate_isin(isin)):
        return reason

    code_raw = values.get("amfi_code")
    amfi_code: int | None = None
    if isinstance(code_raw, int) and not isinstance(code_raw, bool):
        amfi_code = code_raw
    elif (code_text := _text(code_raw)) is not None:
        if not _CODE_RE.fullmatch(code_text):
            return "BAD_AMFI_CODE"
        amfi_code = int(code_text)
    elif code_raw is not None and not isinstance(code_raw, str):
        return "BAD_AMFI_CODE"
    if amfi_code is not None and amfi_code <= 0:
        return "BAD_AMFI_CODE"

    try:
        units = _decimal(values.get("units"))
        cost = _decimal(values.get("cost_basis_inr"))
    except ValueError:
        return "BAD_NUMBER"
    if units is None:
        return "MISSING_UNITS"
    if units <= 0:
        return "UNITS_NOT_POSITIVE"
    if cost is not None and cost < 0:
        return "COST_NEGATIVE"

    d_raw = values.get("purchase_date")
    purchase: date | None
    if d_raw is None or isinstance(d_raw, date):
        purchase = d_raw
    elif (d_text := _text(d_raw)) is None:
        purchase = None
    else:
        try:
            purchase = parse_import_date(d_text)
        except ValueError:
            return "BAD_DATE"

    if isin is None and amfi_code is None:
        flags.append("UNIDENTIFIED")
    return HoldingRow(name, isin, amfi_code, units, cost, purchase, tuple(flags))
